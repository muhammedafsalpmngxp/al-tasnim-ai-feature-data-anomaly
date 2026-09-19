"""Build the Excel workbook: a summary sheet, then one sheet per rule with findings.

STREAMED, NOT ASSEMBLED. Rows are read back from each rule's spool file and appended to a
write-only worksheet one at a time, so a rule with two hundred thousand findings costs the same
memory as one with ten. openpyxl's write_only mode holds no cell objects after they are
written, which is what makes that true rather than merely intended.

openpyxl is imported INSIDE the functions, not at module scope, so a missing dependency
degrades to a clear message about the Excel format alone rather than breaking the whole
application at import time.
"""
from __future__ import annotations

from app.config import settings
from app.observability import get_logger
from app.report import spool as spool_store
from app.report.format import (
    SEVERITY_COLOURS,
    content_widths,
    header,
    number_or_text,
    report_filename,
    sheet_title,
)

log = get_logger()

# The hard limit of the .xlsx format itself, not a policy choice.
XLSX_MAX_ROWS = 1_048_576

_HEADER_FILL = "1F3864"


def _write_summary_sheet(wb, state) -> None:
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.dimensions import ColumnDimension

    ws = wb.create_sheet(title="Summary")
    widths = [14, 52, 24, 12, 16, 16, 12]
    for i, width in enumerate(widths, 1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter] = ColumnDimension(ws, index=letter, width=width)
    ws.freeze_panes = "A2"

    def styled(text, bold=False, colour=None):
        c = WriteOnlyCell(ws, value=text)
        if bold:
            c.font = Font(bold=True, color="FFFFFF" if colour == _HEADER_FILL else "000000")
        if colour:
            c.fill = PatternFill("solid", fgColor=colour)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        return c

    totals = state.get("totals") or {}
    ws.append([styled("Data Quality Report", bold=True)])
    ws.append([f"Run {state.get('run_id', '')}"])
    ws.append([f"Overall score: {state.get('score', 0)} / 100"])
    ws.append([state.get("score_basis", "")])
    ws.append([])
    ws.append([
        f"Probes run: {totals.get('probes_run', 0)}",
        f"With findings: {totals.get('probes_with_findings', 0)}",
        f"Clean: {totals.get('probes_clean', 0)}",
        f"Failed: {totals.get('probes_failed', 0)}",
        f"Records examined: {totals.get('records_examined', 0):,}",
        f"Records flagged: {totals.get('records_flagged', 0):,}",
    ])
    ws.append([])

    if state.get("summary"):
        ws.append([styled("Executive summary", bold=True)])
        for line in str(state["summary"]).splitlines():
            if line.strip():
                ws.append([line.strip()])
        ws.append([])

    ws.append([
        styled(t, bold=True, colour=_HEADER_FILL)
        for t in ("Rule", "What was found", "Category", "Severity",
                  "Records affected", "Of examined", "Share %")
    ])
    for row in state.get("ranked") or []:
        ws.append([
            row["rule_id"],
            row["title"],
            row["category"],
            styled(row["severity"], colour=None),
            row["anomaly_count"],
            row["scope_total"],
            round(float(row["anomaly_pct"]), 2),
        ])

    # Rules still on trial are listed BELOW the findings above, under their own banner and never
    # interleaved with them. Their numbers appear in no total on this sheet: a reader sorting
    # the findings table by "records affected" must not find an unvetted rule at the top of it.
    discovered = state.get("discovered_ranked") or []
    if discovered:
        totals_d = state.get("discovered_totals") or {}
        ws.append([])
        ws.append([styled(
            "ON TRIAL - proposed automatically, accepted for trial, NOT counted in the score "
            "or in any total above",
            bold=True,
        )])
        ws.append([
            styled(t, bold=True, colour=_HEADER_FILL)
            for t in ("Rule", "What was found", "Category", "Severity",
                      "Records affected", "Of examined", "Share %")
        ])
        for row in discovered:
            ws.append([
                row["rule_id"],
                row["title"],
                row["category"],
                styled(row["severity"], colour=None),
                row["anomaly_count"],
                row["scope_total"],
                round(float(row["anomaly_pct"]), 2),
            ])
        ws.append([
            styled(
                f"{totals_d.get('probes_run', 0):,} rule(s) on trial · "
                f"{totals_d.get('records_flagged', 0):,} record(s) flagged · excluded from the "
                f"score and every headline figure",
                bold=True,
            )
        ])

    # Everything NOT running goes in the same workbook, deliberately. A reader who only opens
    # the spreadsheet must still be able to see which checks produced no answer at all.
    gaps = state.get("empty_scope") or []
    not_running = state.get("not_running") or []
    if gaps or not_running:
        ws.append([])
        ws.append([styled("Coverage gaps - these checks did NOT report on their subject",
                          bold=True)])
        for rule_id in gaps:
            ws.append([rule_id, "examined no records at all - this is not a clean result"])
        for item in not_running:
            ws.append([item["rule_id"], f"{item['status']}: {item['reason']}"[:500]])


def _write_rule_sheet(wb, used: set[str], result, rule) -> int:
    """One rule's findings, streamed from its spool file. Returns rows written."""
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.dimensions import ColumnDimension

    title = getattr(rule, "title", result.rule_id)
    ws = wb.create_sheet(title=sheet_title(f"{result.rule_id} {title}", used))

    # MUST come before the first append(): in write_only mode the sheet header is emitted as
    # soon as rows start streaming, so widths set afterwards are silently discarded.
    for i, width in enumerate(content_widths(result.detail_columns, result.detail_rows), 1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter] = ColumnDimension(ws, index=letter, width=width)
    ws.freeze_panes = "A2"

    severity = getattr(rule, "severity", "medium")
    fill = PatternFill("solid", fgColor=SEVERITY_COLOURS.get(severity, _HEADER_FILL))
    head = []
    for name in result.detail_columns:
        c = WriteOnlyCell(ws, value=header(name))
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
        head.append(c)
    ws.append(head)

    body_limit = XLSX_MAX_ROWS - 8  # leave room for the header and the trailing notes
    written = 0
    for row in spool_store.read_rows(result.spool_path):
        if written >= body_limit:
            break
        ws.append([number_or_text(v) for v in row])
        written += 1

    ws.append([])
    ws.append([f"Rule {result.rule_id}: {title}"])
    ws.append([
        f"{result.anomaly_count:,} of {result.scope_total:,} records examined were flagged "
        f"({result.anomaly_pct:.2f}%)."
    ])
    if result.spool_truncated:
        ws.append([
            f"NOTE: this sheet is capped at {settings.export_max_rows:,} rows "
            f"(ANOMALY_EXPORT_MAX_ROWS); the rule found more. Rows are ordered worst first, so "
            f"the most serious are present."
        ])
    if written >= body_limit:
        ws.append([
            f"NOTE: truncated at {body_limit:,} rows - one worksheet cannot hold more than "
            f"{XLSX_MAX_ROWS:,}."
        ])
    for concern in result.concerns:
        ws.append([f"NOTE: {concern}"])
    return written


def build(state, path: str | None = None) -> str:
    """Write the workbook and return its path."""
    from openpyxl import Workbook

    target = path or report_filename(state.get("run_id", "run"), "xlsx")
    # write_only streams rows out instead of holding a cell object per value.
    wb = Workbook(write_only=True)

    _write_summary_sheet(wb, state)

    results = {r.rule_id: r for r in (state.get("results") or [])}
    rules = state.get("rules") or {}
    used: set[str] = {"Summary"}
    sheets = 0
    rows = 0
    # Trial rules get their own sheets too - the affected records are the whole point of running
    # one - but AFTER the established findings, so the workbook reads in order of trust.
    for row in (state.get("ranked") or []) + (state.get("discovered_ranked") or []):
        result = results.get(row["rule_id"])
        if result is None or not result.detail_columns:
            continue
        rows += _write_rule_sheet(wb, used, result, rules.get(result.rule_id))
        sheets += 1

    wb.save(target)
    log.info("report: xlsx written - %d sheet(s), %s row(s) -> %s", sheets, f"{rows:,}", target)
    return target
