"""The one-sheet Summary Excel: every anomaly a run found, one row each.

Columns: rule, anomaly, severity, why it is an anomaly, the tables and columns its query reads,
how many records were affected out of how many examined, and two sample records.

BUILT ON DEMAND FROM WHAT THE RUN ALREADY WROTE. Nothing here touches the database or a model:

  * the findings and counts come from the run's stored results;
  * the two samples are the first two rows of the rule's sheet in the run's full Excel report
    (rows there are ordered worst first, so these are the two most serious);
  * the explanation comes from the rule's own Wrong / Matters text;
  * the tables and columns come from the compiled SQL, read against the cached schema.

So it adds nothing to a detection run, leaves the existing Excel and Word reports exactly as
they were, and works for any past run whose full Excel report is still on disk.
"""
from __future__ import annotations

import re

from app.observability import get_logger
from app.report.format import cell, report_filename

log = get_logger()

SAMPLES_PER_ANOMALY = 2
_CELL_LIMIT = 32_000  # Excel holds at most 32,767 characters in one cell

_NAME = r"(?:\[[^\]]+\]|[A-Za-z_][\w$#@]*)"
_TABLE_REF = re.compile(rf"\b(?:FROM|JOIN)\s+({_NAME}(?:\s*\.\s*{_NAME}){{1,2}})", re.IGNORECASE)
_SECTION = re.compile(r"^\s*([A-Za-z][A-Za-z ]{1,20}):\s*(.*)$")


def _plain(name: str) -> str:
    return ".".join(p.strip().strip("[]").strip() for p in re.split(r"\s*\.\s*", name))


def _why(body: str, title: str) -> str:
    """The rule's Wrong and Matters text, joined - the reason this counts as an anomaly."""
    sections: dict[str, list[str]] = {}
    current = None
    for line in (body or "").splitlines():
        m = _SECTION.match(line)
        if m and not line.startswith((" ", "\t")):
            current = m.group(1).strip().lower()
            sections[current] = [m.group(2).strip()]
        elif current and line.strip() and line.startswith((" ", "\t")):
            sections[current].append(line.strip())
        else:
            current = None
    parts = [" ".join(sections[k]).strip() for k in ("wrong", "matters") if sections.get(k)]
    if parts:
        return "\n".join(parts)
    # The older layout: **What is wrong** / **Why it matters**, each followed by a paragraph.
    text = re.sub(r"\*\*", "", body or "").strip()
    return text[:1500] if text else title


def _tables_and_columns(probe, index) -> tuple[list[str], list[str]]:
    """Tables the probe's SQL reads, and the columns of those tables it mentions."""
    from app.rules.contract import code_of

    if probe is None:
        return [], []
    sql = code_of(f"{probe.summary_sql or ''}\n{probe.detail_sql or ''}")
    known = {name.lower(): name for name in (index.tables if index else {})}
    tables: list[str] = []
    for ref in _TABLE_REF.findall(sql):
        plain = _plain(ref)
        name = known.get(plain.lower(), plain)
        if index is None or plain.lower() in known:
            if name not in tables:
                tables.append(name)
    columns: list[str] = []
    lowered = sql.lower()
    for name in tables:
        table = index.get(name) if index else None
        if table is None:
            continue
        short = name.split(".")[-1]
        used = [
            c.name for c in table.columns
            if f"[{c.name.lower()}]" in lowered
            or re.search(rf"(?<![\w\[]){re.escape(c.name.lower())}(?![\w\]])", lowered)
        ]
        if used:
            columns.append(f"{short}: {', '.join(used)}" if len(tables) > 1 else ", ".join(used))
    return tables, columns


def _samples_by_rule(xlsx_path: str) -> dict[str, list[str]]:
    """The first rows of every rule sheet in the run's full Excel report, as readable text."""
    from openpyxl import load_workbook

    out: dict[str, list[str]] = {}
    if not xlsx_path:
        return out
    try:
        wb = load_workbook(xlsx_path, read_only=True)
    except Exception as exc:  # noqa: BLE001 - no samples is better than no summary
        log.warning("summary xlsx: could not read %s (%s) - samples left blank", xlsx_path, exc)
        return out
    try:
        for sheet in wb.worksheets:
            if sheet.title == "Summary":
                continue
            rule_id = sheet.title.split(" ", 1)[0]
            rows = sheet.iter_rows(max_row=SAMPLES_PER_ANOMALY + 1, values_only=True)
            head = next(rows, None)
            if not head:
                continue
            samples = []
            for row in rows:
                if not row or all(v in (None, "") for v in row):
                    break
                pairs = [
                    f"{h}: {cell(v)}" for h, v in zip(head, row)
                    if h not in (None, "") and v not in (None, "")
                ]
                samples.append("; ".join(pairs)[:_CELL_LIMIT])
            out.setdefault(rule_id, samples)
    finally:
        wb.close()
    return out


def build(run_id: str) -> str:
    """Write the Summary Excel for one run and return its path. Raises LookupError if the run
    has no stored results."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from app.rules import catalog as catalog_store
    from app.rules.expand import expand_families
    from app.rules.loader import load_rules
    from app.runner import history, load_results

    results = load_results(run_id)
    if results is None:
        raise LookupError(f"No stored results for run {run_id!r}")
    row = next((r for r in history() if r.run_id == run_id), None)
    xlsx_path = ((row.report_paths if row else None) or results.get("report_paths") or {}).get("xlsx", "")

    rules, _ = load_rules()
    rules, _ = expand_families(rules)
    rules_by_id = {r.rule_id: r for r in rules}
    catalog = catalog_store.load()
    try:
        from app.rules.schema_index import load_index

        index = load_index()
    except Exception as exc:  # noqa: BLE001 - names from the SQL alone are still useful
        log.warning("summary xlsx: no schema index (%s) - columns left blank", exc)
        index = None
    samples = _samples_by_rule(xlsx_path)

    findings = [
        (f, False) for f in results.get("ranked") or [] if (f.get("anomaly_count") or 0) > 0
    ] + [
        (f, True) for f in results.get("discovered_ranked") or []
        if (f.get("anomaly_count") or 0) > 0
    ]

    columns = [
        "Rule ID", "Anomaly", "Severity", "Why Anomaly", "DB.Table(s)", "Column(s)",
        "Affected Row Count", "Records Examined", "Sample 1", "Sample 2",
    ]
    widths = [13, 40, 10, 60, 48, 48, 12, 12, 70, 70]

    wb = Workbook()
    ws = wb.active
    ws.title = "Anomaly Summary"
    ws.append(columns)
    fill = PatternFill("solid", fgColor="1F3864")
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(vertical="center", wrap_text=True)

    wrap = Alignment(vertical="top", wrap_text=True)
    for finding, on_trial in findings:
        rule_id = finding.get("rule_id", "")
        rule = rules_by_id.get(rule_id)
        title = finding.get("title") or getattr(rule, "title", "") or rule_id
        tables, cols = _tables_and_columns(catalog.get(rule_id), index)
        found = (samples.get(rule_id) or [])[:SAMPLES_PER_ANOMALY]
        found += [""] * (SAMPLES_PER_ANOMALY - len(found))
        ws.append([
            rule_id,
            title + (" (rule on trial)" if on_trial else ""),
            str(finding.get("severity", "")).title(),
            _why(getattr(rule, "body", ""), title)[:_CELL_LIMIT],
            "\n".join(tables),
            "\n".join(cols)[:_CELL_LIMIT],
            finding.get("anomaly_count") or 0,
            finding.get("scope_total") or 0,
            *found,
        ])
        for c in ws[ws.max_row]:
            c.alignment = wrap

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(ws.max_row, 1)}"

    target = report_filename(run_id, "xlsx", prefix="data-quality-summary")
    wb.save(target)
    log.info("report: summary xlsx written - %d anomaly row(s) -> %s", len(findings), target)
    return target

